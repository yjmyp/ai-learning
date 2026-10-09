# -*- coding: utf-8 -*-
"""
doc_io · 简历/资料读取（PDF / Word / 文本 / 图片 OCR）
======================================================
支持格式与实现方式：
  .pdf            pypdf 逐页抽文字（扫描版 PDF 抽不出文字，会提示）
  .docx           python-docx 逐段抽文字
  .txt / .md      多编码尝试（utf-8-sig → utf-8 → gb18030）
  .png/.jpg/.jpeg/.webp/.bmp   RapidOCR 离线识别（中英混排）

设计原则：**识别失败要说清楚失败原因和替代方案**，不要静默返回空字符串。
"""
import io

TEXT_EXT = {".txt", ".md", ".markdown"}
DOC_EXT = {".docx"}
PDF_EXT = {".pdf"}
IMG_EXT = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}
ALL_EXT = TEXT_EXT | DOC_EXT | PDF_EXT | IMG_EXT


def _read_text(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return data.decode(enc)
        except Exception:
            continue
    return data.decode("utf-8", errors="ignore")


def _read_pdf(data: bytes) -> tuple:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    pages = [(p.extract_text() or "") for p in reader.pages]
    text = "\n".join(pages).strip()
    if len(text) < 40:
        return text, ("这份 PDF 里几乎抽不到文字，可能是扫描/图片版。"
                      "可以改用截图上传（走 OCR），或直接把文字粘进来。")
    return text, ""


def _read_docx(data: bytes) -> tuple:
    from docx import Document
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip(), ""


def _read_image(data: bytes) -> tuple:
    """图片 → 文字。两条路，谁可用走谁：

      1. RapidOCR（本地常用）：中英混排准、离线，但依赖 opencv；
         实测把它装到 Streamlit Cloud 会把 opencv 拖进来导致构建失败、应用起不来。
      2. pytesseract + 系统 tesseract（云端走这条）：只依赖 apt 包，装得快、不碰 opencv。
    两条都按"懒导入"处理——**缺组件只让这一个功能报错，不影响应用启动**（这条是真踩过教训）。
    """
    try:
        from rapidocr_onnxruntime import RapidOCR
    except Exception:
        return _read_image_tesseract(data)
    import tempfile
    from pathlib import Path
    tmp = Path(tempfile.mkdtemp()) / "upload.png"
    tmp.write_bytes(data)
    try:
        engine = RapidOCR()
        result, _ = engine(str(tmp))
    except Exception as e:
        # RapidOCR 装了但跑不起来（例如缺 libGL）→ 退到 tesseract
        return _read_image_tesseract(data, note=f"RapidOCR 不可用（{type(e).__name__}）")
    lines = [x[1] for x in (result or [])]
    text = "\n".join(lines).strip()
    warn = "" if len(text) > 20 else "OCR 只识别出很少文字，图片可能太模糊或字太小。"
    return text, warn


def _read_image_tesseract(data: bytes, note: str = "") -> tuple:
    """tesseract 兜底（Streamlit Cloud 用这条：apt 装 tesseract-ocr + chi_sim）。"""
    try:
        import pytesseract
        from PIL import Image
        import io
        img = Image.open(io.BytesIO(data))
        text = pytesseract.image_to_string(img, lang="chi_sim+eng").strip()
        if len(text) > 20:
            return text, note
        return text, (note + " " if note else "") + "OCR 只识别出很少文字，图片可能太模糊或字太小。"
    except Exception as e:
        return "", ("没装 OCR 组件（云端需要 packages.txt 里的 tesseract-ocr + tesseract-ocr-chi-sim，"
                    "并安装 pytesseract）。本机可跑：python -m pip install rapidocr-onnxruntime。"
                    "当前错误：%s" % type(e).__name__)


def read_any(data: bytes, filename: str) -> dict:
    """统一入口。返回 {text, method, warning}"""
    ext = ("." + filename.rsplit(".", 1)[-1].lower()) if "." in filename else ""
    if ext in TEXT_EXT:
        return {"text": _read_text(data), "method": "文本直读", "warning": ""}
    if ext in PDF_EXT:
        t, w = _read_pdf(data)
        return {"text": t, "method": "PDF 抽文字", "warning": w}
    if ext in DOC_EXT:
        t, w = _read_docx(data)
        return {"text": t, "method": "Word 抽文字", "warning": w}
    if ext in IMG_EXT:
        t, w = _read_image(data)
        return {"text": t, "method": "图片 OCR", "warning": w}
    return {"text": "", "method": "不支持",
            "warning": f"不支持的文件类型 {ext or '（无扩展名）'}。"
                       f"支持：{'、'.join(sorted(ALL_EXT))}"}
